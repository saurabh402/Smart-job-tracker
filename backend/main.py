import os
import re
from contextlib import asynccontextmanager
from datetime import date
from typing import Optional

import psycopg
from psycopg.rows import dict_row
from dotenv import load_dotenv

from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from pypdf import PdfReader
from docx import Document


# =========================================================
# CONFIGURATION
# =========================================================

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))

DATABASE_URL = os.getenv("DATABASE_URL")


def get_connection():
    if not DATABASE_URL:
        raise RuntimeError(
            "DATABASE_URL missing hai. backend/.env file check karein."
        )

    return psycopg.connect(
        DATABASE_URL,
        row_factory=dict_row
    )


# =========================================================
# DATABASE SETUP
# Existing data delete nahi hoga.
# =========================================================

@asynccontextmanager
async def lifespan(app: FastAPI):
    with get_connection() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS applications (
                id SERIAL PRIMARY KEY,
                company TEXT NOT NULL,
                role TEXT NOT NULL,
                location TEXT DEFAULT '',
                applied_date DATE,
                status TEXT DEFAULT 'Applied',
                job_type TEXT DEFAULT '',
                job_url TEXT DEFAULT '',
                notes TEXT DEFAULT '',
                follow_up_date DATE,
                follow_up_done BOOLEAN DEFAULT FALSE,
                interview_date DATE,
                interview_time TEXT DEFAULT '',
                interview_round TEXT DEFAULT '',
                interview_mode TEXT DEFAULT '',
                interviewer TEXT DEFAULT '',
                interview_notes TEXT DEFAULT '',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        conn.execute("""
            ALTER TABLE applications
            ADD COLUMN IF NOT EXISTS follow_up_date DATE
        """)

        conn.execute("""
            ALTER TABLE applications
            ADD COLUMN IF NOT EXISTS follow_up_done BOOLEAN DEFAULT FALSE
        """)

    yield


# =========================================================
# FASTAPI APP
# =========================================================

app = FastAPI(
    title="Smart Job Tracker API",
    lifespan=lifespan
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"]
)


# =========================================================
# APPLICATION MODEL
# =========================================================

class Application(BaseModel):
    company: str
    role: str
    location: Optional[str] = ""
    appliedDate: Optional[date] = None
    status: Optional[str] = "Applied"
    jobType: Optional[str] = ""
    jobUrl: Optional[str] = ""
    notes: Optional[str] = ""
    followUpDate: Optional[date] = None
    followUpDone: Optional[bool] = False
    interviewDate: Optional[date] = None
    interviewTime: Optional[str] = ""
    interviewRound: Optional[str] = ""
    interviewMode: Optional[str] = ""
    interviewer: Optional[str] = ""
    interviewNotes: Optional[str] = ""


def application_dict(row):
    return {
        "id": row["id"],
        "company": row["company"],
        "role": row["role"],
        "location": row["location"] or "",
        "appliedDate": (
            row["applied_date"].isoformat()
            if row["applied_date"] else ""
        ),
        "status": row["status"] or "Applied",
        "jobType": row["job_type"] or "",
        "jobUrl": row["job_url"] or "",
        "notes": row["notes"] or "",
        "followUpDate": (
            row["follow_up_date"].isoformat()
            if row["follow_up_date"] else ""
        ),
        "followUpDone": row["follow_up_done"] or False,
        "interviewDate": (
            row["interview_date"].isoformat()
            if row["interview_date"] else ""
        ),
        "interviewTime": row["interview_time"] or "",
        "interviewRound": row["interview_round"] or "",
        "interviewMode": row["interview_mode"] or "",
        "interviewer": row["interviewer"] or "",
        "interviewNotes": row["interview_notes"] or "",
        "createdAt": (
            row["created_at"].isoformat()
            if row.get("created_at") else ""
        )
    }


def model_to_db(data: Application):
    return {
        "company": data.company,
        "role": data.role,
        "location": data.location or "",
        "applied_date": data.appliedDate,
        "status": data.status or "Applied",
        "job_type": data.jobType or "",
        "job_url": data.jobUrl or "",
        "notes": data.notes or "",
        "follow_up_date": data.followUpDate,
        "follow_up_done": data.followUpDone or False,
        "interview_date": data.interviewDate,
        "interview_time": data.interviewTime or "",
        "interview_round": data.interviewRound or "",
        "interview_mode": data.interviewMode or "",
        "interviewer": data.interviewer or "",
        "interview_notes": data.interviewNotes or ""
    }


# =========================================================
# EXISTING APPLICATION ENDPOINTS
# =========================================================

@app.get("/")
def home():
    return {
        "message": "Smart Job Tracker API is running."
    }


@app.get("/api/health")
def health():
    try:
        with get_connection() as conn:
            conn.execute("SELECT 1")

        return {
            "status": "ok",
            "database": "connected"
        }

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=f"Database connection error: {str(error)}"
        )


@app.get("/api/applications")
def get_applications():
    try:
        with get_connection() as conn:
            rows = conn.execute("""
                SELECT *
                FROM applications
                ORDER BY created_at DESC
            """).fetchall()

        return [application_dict(row) for row in rows]

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=str(error)
        )


@app.post("/api/applications", status_code=201)
def add_application(data: Application):
    try:
        values = model_to_db(data)

        with get_connection() as conn:
            row = conn.execute("""
                INSERT INTO applications (
                    company, role, location, applied_date, status,
                    job_type, job_url, notes, follow_up_date,
                    follow_up_done, interview_date, interview_time,
                    interview_round, interview_mode, interviewer,
                    interview_notes
                )
                VALUES (
                    %(company)s, %(role)s, %(location)s, %(applied_date)s,
                    %(status)s, %(job_type)s, %(job_url)s, %(notes)s,
                    %(follow_up_date)s, %(follow_up_done)s,
                    %(interview_date)s, %(interview_time)s,
                    %(interview_round)s, %(interview_mode)s,
                    %(interviewer)s, %(interview_notes)s
                )
                RETURNING *
            """, values).fetchone()

        return application_dict(row)

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=str(error)
        )


@app.put("/api/applications/{application_id}")
def update_application(application_id: int, data: Application):
    try:
        values = model_to_db(data)
        values["id"] = application_id

        with get_connection() as conn:
            row = conn.execute("""
                UPDATE applications
                SET
                    company = %(company)s,
                    role = %(role)s,
                    location = %(location)s,
                    applied_date = %(applied_date)s,
                    status = %(status)s,
                    job_type = %(job_type)s,
                    job_url = %(job_url)s,
                    notes = %(notes)s,
                    follow_up_date = %(follow_up_date)s,
                    follow_up_done = %(follow_up_done)s,
                    interview_date = %(interview_date)s,
                    interview_time = %(interview_time)s,
                    interview_round = %(interview_round)s,
                    interview_mode = %(interview_mode)s,
                    interviewer = %(interviewer)s,
                    interview_notes = %(interview_notes)s
                WHERE id = %(id)s
                RETURNING *
            """, values).fetchone()

        if not row:
            raise HTTPException(
                status_code=404,
                detail="Application nahi mili."
            )

        return application_dict(row)

    except HTTPException:
        raise

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=str(error)
        )


@app.delete("/api/applications/{application_id}")
def delete_application(application_id: int):
    try:
        with get_connection() as conn:
            row = conn.execute("""
                DELETE FROM applications
                WHERE id = %s
                RETURNING id
            """, (application_id,)).fetchone()

        if not row:
            raise HTTPException(
                status_code=404,
                detail="Application nahi mili."
            )

        return {
            "message": "Application deleted successfully.",
            "id": application_id
        }

    except HTTPException:
        raise

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=str(error)
        )


# =========================================================
# RESUME ANALYZER - SKILL LIST
# Yeh transparent, rules-based estimate hai.
# Actual proprietary ATS score nahi hai.
# =========================================================

SKILL_KEYWORDS = [
    "python",
    "sql",
    "excel",
    "power bi",
    "tableau",
    "pandas",
    "numpy",
    "matplotlib",
    "seaborn",
    "statistics",
    "data analysis",
    "data analytics",
    "data visualization",
    "machine learning",
    "deep learning",
    "artificial intelligence",
    "communication",
    "problem solving",
    "critical thinking",
    "project management",
    "business analysis",
    "business intelligence",
    "postgresql",
    "mysql",
    "mongodb",
    "database",
    "html",
    "css",
    "javascript",
    "react",
    "node.js",
    "fastapi",
    "flask",
    "django",
    "git",
    "github",
    "api",
    "rest api",
    "statistics",
    "probability",
    "forecasting",
    "etl",
    "data cleaning",
    "data modeling",
    "reporting",
    "dashboard",
    "pivot table",
    "vlookup",
    "power query",
    "financial modeling",
    "communication skills",
    "teamwork",
    "leadership",
    "research",
    "time management"
]


def clean_text(text: str) -> str:
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    extension = os.path.splitext(filename.lower())[1]

    try:
        if extension == ".pdf":
            import io

            reader = PdfReader(io.BytesIO(file_bytes))

            text_parts = []

            for page in reader.pages:
                text_parts.append(page.extract_text() or "")

            return "\n".join(text_parts)

        elif extension == ".docx":
            import io

            document = Document(io.BytesIO(file_bytes))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            table_text = []

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        table_text.append(cell.text)

            return "\n".join(paragraphs + table_text)

        elif extension == ".txt":
            return file_bytes.decode("utf-8", errors="ignore")

        else:
            raise HTTPException(
                status_code=400,
                detail="Sirf PDF, DOCX aur TXT files supported hain."
            )

    except HTTPException:
        raise

    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail=f"Resume read nahi ho paya: {str(error)}"
        )


def find_skills(text: str):
    normalized = clean_text(text)
    found = []

    for skill in SKILL_KEYWORDS:
        pattern = r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])"

        if re.search(pattern, normalized):
            found.append(skill)

    return sorted(set(found))


def analyze_resume(resume_text: str, job_description: str):
    if not resume_text.strip():
        raise HTTPException(
            status_code=400,
            detail=(
                "Resume se readable text nahi mila. "
                "Agar PDF scanned image hai, to text-based PDF use karein."
            )
        )

    if not job_description.strip():
        raise HTTPException(
            status_code=400,
            detail="Job description khaali hai."
        )

    resume_skills = set(find_skills(resume_text))
    job_skills = set(find_skills(job_description))

    if not job_skills:
        raise HTTPException(
            status_code=400,
            detail=(
                "Job description mein known skills nahi mili. "
                "Job description mein required skills jaise Python, SQL, Excel likhein."
            )
        )

    matched_skills = sorted(resume_skills.intersection(job_skills))
    missing_skills = sorted(job_skills - resume_skills)

    score = round(
        (len(matched_skills) / len(job_skills)) * 100
    )

    suggestions = []

    if missing_skills:
        suggestions.append(
            "In missing skills ko seekhein ya apne relevant experience "
            "ke saath resume mein add karein, agar sach mein aata hai: "
            + ", ".join(missing_skills[:10])
        )

    if not resume_skills:
        suggestions.append(
            "Resume mein relevant technical skills aur project experience "
            "clear headings ke saath mention karein."
        )

    if len(resume_text.split()) < 100:
        suggestions.append(
            "Resume ka text kaafi chhota lag raha hai. Apne projects, "
            "responsibilities, education aur measurable results add karein."
        )

    if score >= 80:
        suggestions.append(
            "Achha skill overlap hai. Resume ko job description ke hisaab "
            "se tailor karein aur apne projects ke results clearly dikhayein."
        )
    elif score >= 50:
        suggestions.append(
            "Kuch important skills match ho rahi hain. Missing skills par "
            "kaam karein aur relevant projects ko resume mein highlight karein."
        )
    else:
        suggestions.append(
            "Is role ke liye skill overlap kam hai. Job description ki "
            "priority skills par practice aur projects banane par focus karein."
        )

    return {
        "matchScore": score,
        "matchedSkills": matched_skills,
        "missingSkills": missing_skills,
        "matchedCount": len(matched_skills),
        "totalJobSkills": len(job_skills),
        "suggestions": suggestions,
        "scoreNote": (
            "Yeh listed skills ke overlap par based estimated score hai. "
            "Yeh kisi company ka actual ATS score, hiring probability ya "
            "job selection ki guarantee nahi hai."
        )
    }


# =========================================================
# RESUME ANALYZER API
# POST /api/resume/analyze
# =========================================================

@app.post("/api/resume/analyze")
async def resume_analyzer(
    resume: UploadFile = File(...),
    job_description: str = Form(...)
):
    allowed_extensions = {".pdf", ".docx", ".txt"}

    filename = resume.filename or "resume.txt"
    extension = os.path.splitext(filename.lower())[1]

    if extension not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail="Resume PDF, DOCX ya TXT format mein upload karein."
        )

    file_bytes = await resume.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="Uploaded resume khaali hai."
        )

    # Maximum upload size: 5 MB
    if len(file_bytes) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail="Resume file 5 MB se chhoti honi chahiye."
        )

    resume_text = extract_text_from_file(filename, file_bytes)

    return analyze_resume(
        resume_text=resume_text,
        job_description=job_description
    )